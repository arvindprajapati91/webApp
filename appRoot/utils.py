from io import BytesIO
import random
import string
from datetime import datetime
from django.conf import settings
from django.db import transaction
from django.contrib.auth.hashers import make_password
from django.contrib.auth import get_user_model
from functools import wraps
import sys, os, traceback
from django.http import HttpResponse,JsonResponse
import openpyxl as xl
from openpyxl.styles import Font,Border, Side
from django.db.models import F
import zipfile
from django.http import HttpResponse
from django.template.loader import get_template
from xhtml2pdf import pisa
from appRoot.settings import BASE_DIR, STATIC_ROOT
from .views import *
from django.core.mail import EmailMessage
from django.http import HttpResponse
import yaml
from yaml.loader import SafeLoader
from administration.auditTrail.models import Audit_Trail
from administration.countryMaster.models import Country_Master
from administration.departmentMaster.models import Department_Master
from administration.locationMaster.models import Location_Master
from administration.processMaster.models import Process_Master
from administration.statusMaster.models import Status_Master
from administration.organisation.models import Organisation#, Organisation_Ladder
from administration.businessLine.models import Business_Line
from administration.seriesMaster.models import Series_Master as SM
from auditlog.models import *
from django.contrib.contenttypes.models import ContentType

User = get_user_model()

def random_string_generator(size=10,chars=string.ascii_lowercase + string.digits):
    return ''.join(random.choice(chars) for _ in range(size))

def session_expired():
    data = {
        "error":"session expired"
	}
    return JsonResponse(data)

#Error Handler Function
def catch_exceptions(function):
    @wraps(function)
    def decorator(request,*args, **kwargs):
        userName = request.user
        try:
            return function(request,*args, **kwargs)
        except Exception as e:
            error_log_file(e,userName)
    return decorator

def error_log_file(e,userName):
    exc_type, exc_obj, exc_tb = sys.exc_info()
    error = "E=%s, F=%s, L=%s" % (str(e), traceback.extract_tb(exc_tb)[-1][0], traceback.extract_tb(exc_tb)[-1][1])
    dirBase = settings.BASE_DIR
    date_time_str = datetime.now().strftime("%d-%b-%Y")
    file = dirBase/ f'./logs/errorLog_{date_time_str}.txt'
    f = open(file, "a")
    f.write("\n")
    f.write("Username : "+str(userName)+", TimeStamp : "+ str(datetime.now())+"\n")
    f.write(error+"\n")
    f.write("---------------------------------------------------------------------------------------------------------------------------------------")
    f.close()
    exp_error = error.split(",")
    err = ""
    for i in exp_error:
        if "F=" in i:
            pass
        else:
            if err == "":
                err = err+i
            else:
                err = err+" "+i
    data = {
            "error_msg": err+"\n"
        }
    return JsonResponse(data,safe=False)

# Create Update Function
@transaction.atomic
@catch_exceptions
def objSave(request,obj,body_data,obj_id,field_names):
    modelName = obj._meta.model.__name__
    model = obj._meta.model
    request.session['exception_message'] = ""
    modelFields = [f.name for f in model._meta.get_fields()]
    if modelName == "User":
        obj.process = Process_Master.objects.filter(process_desc="User Master").first()
        for key, value in body_data.items():
            if key == 'is_lock':
                if value == '1':
                    setattr(obj, key, True)
                else:
                    setattr(obj, key, False)

            elif key == "random_password" :
                if value != '':
                    obj.password=make_password(value)
            elif key == "username":
                setattr(obj, key, value.casefold())
            elif key == "email":
                setattr(obj, key, value.casefold())
            elif key == "managerSelected":
                Manager = User.objects.filter(full_name=value).first()
                obj.reporting_manager = Manager
            elif key == "orgSelected":
                Org = Organisation.objects.filter(org_desc=value).first()
                obj.organisation = Org
            elif key == "blSelected":
                BL = Business_Line.objects.filter(bl_desc=value).first()
                obj.business_line = BL
            elif key == "depSelected":
                Dep = Department_Master.objects.filter(dep_desc=value).first()
                obj.department = Dep
            elif key == "locSelected":
                Loc = Location_Master.objects.filter(loc_desc=value).first()
                obj.location = Loc

            # Role Access
            elif key == "role_ids":
                role_access_list = []
                role_list = Role_Master.objects.values_list('id', flat=True).filter(id__in=value).distinct()
                if role_list:
                    for i in role_list:
                        role_access_list.append(i)
                obj.role_access.clear()
                obj.role_access.add(* role_access_list)

            # Org Access
            elif key == "org_ids":
                org_access_list = []
                _list = Organisation.objects.values_list('id', flat=True).filter(id__in=value).distinct()
                if _list:
                    for i in _list:
                        org_access_list.append(i)
                obj.org_access.clear()
                obj.org_access.add(* org_access_list)

            # BL Access
            elif key == "bl_ids":
                bl_access_list = []
                _list = Business_Line.objects.values_list('id', flat=True).filter(id__in=value).distinct()
                if _list:
                    for i in _list:
                        bl_access_list.append(i)
                obj.bl_access.clear()
                obj.bl_access.add(* bl_access_list)

            # Dep Access
            elif key == "dep_ids":
                dep_access_list = []
                _list = Department_Master.objects.values_list('id', flat=True).filter(id__in=value).distinct()
                if _list:
                    for i in _list:
                        dep_access_list.append(i)
                obj.dep_access.clear()
                obj.dep_access.add(* dep_access_list)

            # Loc Access
            elif key == "loc_ids":
                loc_access_list = []
                _list = Location_Master.objects.values_list('id', flat=True).filter(id__in=value).distinct()
                if _list:
                    for i in _list:
                        loc_access_list.append(i)
                obj.loc_access.clear()
                obj.loc_access.add(* loc_access_list)

            elif key == "status":
                if value == "1":
                    setattr(obj, key, True)
                else:
                    setattr(obj, key, False)
            else:
                if key in field_names:
                    setattr(obj, key, value)
    elif modelName == "Organisation":
        for key, value in body_data.items():
            if key == "countrySelected":
                cm = Country_Master.objects.filter(flag=value).first()
                obj.country_id = cm.id
            if key == "stateMasterSelected":
                sm = State_Master.objects.filter(state_code=value).first()
                obj.state_id = sm.id
            elif key == "status" :
                if value == "1":
                    setattr(obj, key, True)
                else:
                    setattr(obj, key, False)
            else:
                if key in field_names:
                    setattr(obj, key, value)
    else:
        for key, value in body_data.items():
            if key == "process" :
                obj.process = Process_Master.objects.filter(process_desc=value).first()
            elif key == "status_master" :
                obj.status = Status_Master.objects.filter(status_desc=value).first()
            elif modelName == "Process_Master" and key == "process_type":
                setattr(obj, "type", value)
            elif key == "countrySelected":
                cm = Country_Master.objects.filter(flag=value).first()
                obj.country_id = cm.id
            elif key == "orgSelected":
                org = Organisation.objects.filter(org_code=value).first()
                obj.organisation_id = org.id
            elif key == "locSelected":
                loc = Location_Master.objects.filter(loc_code=value).first()
                obj.location = loc
            elif key == "depSelected":
                dep = Department_Master.objects.filter(dep_code=value).first()
                obj.department = dep
            elif key == "status" :
                if value == "1":
                    setattr(obj, key, True)
                else:
                    setattr(obj, key, False)
            else:
                if key in field_names:
                    if value == "child" and key == 'is_parent':
                        value = False
                    elif value == "parent" and key == 'is_parent':
                        value = True
                    else:
                        value = value
                    setattr(obj, key, value)
    if obj_id == None:
        obj.created_by = request.user
        obj.slug=str(random_string_generator())+str(obj.id)
    if "organisation" in modelFields and obj.organisation_id is None:
        org = request.user.organisation
        obj.organisation = org
    if modelName == "Inward_Entry":
        obj.status = "Created"
        obj.process = Process_Master.objects.filter(process_desc="Inward Entry").first()
    obj.updated_by = request.user
    obj.save()

def approvalSave(request, process_name, modelName):
    obj_id = request.POST.get("obj_id")
    comments = request.POST.get("comments")
    action = request.POST.get("action")
    process = Process_Master.objects.filter(process_desc=process_name).first()
    qs = modelName.objects.filter(id=obj_id).first()
    user_rl = User.objects.filter(id=request.user.id).values_list('role_access__id', flat=True)
    if action == "approveBtn" :
        processLadrQS = Process_Ladder.objects.filter(unq_id=obj_id).filter(process=process).first()
        if processLadrQS.role.id in user_rl:
            processLadrQS.delete()
            audit_trail_save(request=request, obj_id=obj_id, process=process,action_desc="Approve", changes="", comments=comments)
            processLadrQS = Process_Ladder.objects.filter(unq_id=obj_id).filter(process=process).first()
            if processLadrQS is None:
                seriesMaster = SM.objects.filter(process=process).filter(status__status_desc="Approved").first()
                nextNo = seriesMaster.series_no+1
                doc_no = qs.type+"-NOM-"+str(nextNo).zfill(10)
                qs.doc_no=doc_no
                status = Status_Master.objects.filter(status_desc="Approved").first()
                qs.status=status
                qs.save()
                seriesMaster.series_no=nextNo
                seriesMaster.save()
    elif action == "moreInfoBtn" :
        processLadrQS = Process_Ladder.objects.filter(unq_id=obj_id).filter(process=process).all()
        processLadrQS.delete()
        audit_trail_save(request=request, obj_id=obj_id, process=process,action_desc="More Info", changes="", comments=comments)
        status = Status_Master.objects.filter(status_desc="More Info").first()
        qs.status=status
        qs.save()
    elif action == "rejectBtn" :
        processLadrQS = Process_Ladder.objects.filter(unq_id=obj_id).filter(process=process).all()
        processLadrQS.delete()
        audit_trail_save(request=request, obj_id=obj_id, process=process,action_desc="Reject", changes="", comments=comments)
        status = Status_Master.objects.filter(status_desc="Rejected").first()
        qs.status=status
        qs.save()

    data = {
        "status":"Ok"
    }
    return data

def audit_trail_save(request,obj_id,process,action_desc,changes,comments):
    audit_trail = Audit_Trail(unq_id=obj_id, process=process, action_desc=action_desc, action_by=request.user, changes=changes,comments=comments)
    audit_trail.save()


def adminFormFeilds(ModelName):
    field_names = []
    for f in ModelName._meta.get_fields():
        if f.many_to_many is True or f.one_to_many is True or f.one_to_one is True or f.many_to_one is True or f.many_to_one is True:
            pass
        else:
            field_names.append(f.name)
    return field_names


def download_excel_data(request,modelName,sheetName,fields):
    model_criteria = {}
    parameters = request.GET.dict()
    if len(parameters) > 0:
        for key, value in parameters.items():
            model_criteria.update({key+'__icontains':value})
        qs = modelName.objects.select_related().filter(**model_criteria).order_by('name').all().values().annotate(created_by=F('created_by__full_name'),updated_by=F('updated_by__full_name'))
    else:
        qs = modelName.objects.select_related().order_by('name').all().values().annotate(created_by=F('created_by__full_name'),updated_by=F('updated_by__full_name'))

    wb = xl.Workbook()
    ws =  wb.active
    ws.append(fields)
    for i in qs.values():
        v = []
        for f in fields:
            field_name = f
            value = i[f]
            if field_name == 'status':
                if value is True:
                    value = 'Active'
                else:
                    value = 'InActive'
            if value == 'NULL':
                value = ''
            if value is True:
                value = 'Yes'
            if value is False:
                value = 'No'
            v.append(value)
        ws.append(v)
        for rows in ws.iter_rows(min_row=1, max_row=1):
            for cell in rows:
                cell.font = Font(bold=True)
        ws.title = sheetName
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(cell.value)
                except:
                    pass
            adjusted_width = (max_length + 2) * 1.30
            ws.column_dimensions[column_letter].width = adjusted_width
            ws.column_dimensions[column_letter].font = Font(size=10)
    filename = "simple.xlsx"
    response = HttpResponse(content_type='application/ms-excel')
    response["Content-Disposition"] = "attachment; filename=%s" % filename
    wb.save(response)
    return response


def render_to_pdf(template_src, context_dict={}):
    template = get_template(template_src)
    html  = template.render(context_dict)
    result = BytesIO()
    pdf = pisa.pisaDocument(BytesIO(html.encode("ISO-8859-1")), result)
    if pdf.err:
        return HttpResponse("Invalid PDF", status_code=400, content_type='text/plain')
    return HttpResponse(result.getvalue(), content_type='application/pdf')


def generate_zip(files):
    mem_zip = BytesIO()
    with zipfile.ZipFile(mem_zip, mode="w",compression=zipfile.ZIP_DEFLATED) as zf:
        for file in files:
            fileName= str(file[0])
            if ".pdf" in fileName:
                content = file[1].content
                zf.writestr(fileName,content)

    mem_zip.flush()
    return mem_zip.getvalue()


def sendEmail(to_email_id,from_email_id,email_message,subject,attachment):
    html_content = """<html>
        <body>
            <p>Hi,</p>
            <p>"""+email_message+"""</p>
            <p>Regards,
                <br>System Admin
            </p>
            <p></p>
        </body>
        </html>"""

    email = EmailMessage(subject, html_content,from_email_id, [to_email_id])
    email.content_subtype = "html"
    email.attach(subject+'.xlsm', attachment.getvalue(), 'application/ms-excel')
    res = email.send()
    return HttpResponse('%s'%res)


def autoAdjustColumns(sheet):
    for column in sheet.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            cell.font = Font(size=9)
            try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(cell.value)
            except:
                    pass
        adjusted_width = (max_length)
        sheet.column_dimensions[column_letter].width = adjusted_width
        sheet.column_dimensions[column_letter].font = Font(size=9)

def auditTrailSave(request,action_desc,modelName,obj,record_no,process):
    content_type = ContentType.objects.get(model=modelName._meta.model_name)
    changeHistory = LogEntry.objects.filter(object_id=obj.id).filter(content_type_id=content_type.id)
    if changeHistory.exists():
        latestChange = changeHistory.latest('timestamp')
        latestChange = latestChange.changes_str
        if "None" in latestChange:
            latestChange = ""
        else:
            latestChange = latestChange
            auditTrail = Audit_Trail.objects.filter(record_no=record_no,process=process).last()
            if auditTrail is not None and latestChange == auditTrail.changes:
                latestChange = ""
            else:
                latestChange = latestChange
    else:
        latestChange = ""
    latestChange=latestChange
    audit_trail = Audit_Trail(table_name=modelName._meta.model_name,record_no=record_no, unq_id=obj.id, process=process, action_desc=action_desc, action_by=request.user, changes=latestChange,comments="")
    audit_trail.save()